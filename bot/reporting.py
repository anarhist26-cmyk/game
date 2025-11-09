"""Generate XLSX reports from aggregated Bnovo statistics."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .bnovo_client import DailyStat


HEADER_STYLE = NamedStyle(
    name="bnovo_header",
    font=Font(bold=True),
    alignment=Alignment(horizontal="center"),
)


def render_report(
    stats: Sequence[DailyStat],
    *,
    template_path: Optional[Path] = None,
    include_expenses: bool = False,
) -> bytes:
    """Return an XLSX file populated with the provided statistics."""

    workbook, worksheet, start_row = _load_template(template_path)
    current_row = start_row

    for stat in stats:
        worksheet.cell(row=current_row, column=1, value=stat.day.strftime("%d.%m.%Y"))
        worksheet.cell(row=current_row, column=2, value=stat.occupancy)
        worksheet.cell(row=current_row, column=3, value=round(stat.base_amount, 2))
        worksheet.cell(row=current_row, column=4, value=round(stat.extra_amount, 2))
        worksheet.cell(row=current_row, column=5, value=round(stat.total_amount, 2))

        worksheet.cell(row=current_row, column=3).number_format = "#,##0.00"
        worksheet.cell(row=current_row, column=4).number_format = "#,##0.00"
        worksheet.cell(row=current_row, column=5).number_format = "#,##0.00"
        current_row += 1

    summary_row = current_row
    worksheet.cell(row=summary_row, column=1, value="ИТОГО")
    worksheet.cell(
        row=summary_row,
        column=2,
        value=f"=SUM(B{start_row}:B{current_row - 1})",
    )
    worksheet.cell(
        row=summary_row,
        column=3,
        value=f"=SUM(C{start_row}:C{current_row - 1})",
    )
    worksheet.cell(
        row=summary_row,
        column=4,
        value=f"=SUM(D{start_row}:D{current_row - 1})",
    )
    worksheet.cell(
        row=summary_row,
        column=5,
        value=f"=SUM(E{start_row}:E{current_row - 1})",
    )

    worksheet.cell(row=summary_row, column=3).number_format = "#,##0.00"
    worksheet.cell(row=summary_row, column=4).number_format = "#,##0.00"
    worksheet.cell(row=summary_row, column=5).number_format = "#,##0.00"

    if include_expenses:
        worksheet.cell(row=summary_row + 2, column=1, value="Расход")
        worksheet.cell(row=summary_row + 3, column=1, value="интернет")
        worksheet.cell(row=summary_row + 4, column=1, value="реклама")
        worksheet.cell(
            row=summary_row + 6,
            column=1,
            value="ИТОГО",
        )
        worksheet.cell(
            row=summary_row + 6,
            column=5,
            value=f"=E{summary_row}-SUM(E{summary_row + 2}:E{summary_row + 5})",
        )
        worksheet.cell(row=summary_row + 6, column=5).number_format = "#,##0.00"

    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def _load_template(
    template_path: Optional[Path],
) -> tuple[Workbook, Worksheet, int]:
    if template_path and template_path.exists():
        workbook = load_workbook(template_path)
        worksheet = workbook.active
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Отчет"
        worksheet.append(["Дата", "Броней", "Сумма", "Доп. место", "Итого"])
        try:
            if HEADER_STYLE.name not in workbook.named_styles:
                workbook.add_named_style(HEADER_STYLE)
        except ValueError:
            # Style already registered
            pass
        for column, width in enumerate([15, 12, 18, 18, 18], start=1):
            worksheet.column_dimensions[get_column_letter(column)].width = width
        for column in range(1, 6):
            cell = worksheet.cell(row=1, column=column)
            cell.style = HEADER_STYLE
    start_row = _detect_start_row(worksheet)
    return workbook, worksheet, start_row


def _detect_start_row(worksheet: Worksheet) -> int:
    max_row = worksheet.max_row
    for row in range(2, max_row + 5):
        if all(
            worksheet.cell(row=row, column=col).value in (None, "")
            for col in range(1, 6)
        ):
            return row
    return max_row + 1


__all__ = ["render_report"]
